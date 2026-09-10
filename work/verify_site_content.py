from html import escape
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sources = {
    '2024': Path(r'C:\Users\noahk\Downloads\2024 Fantasy Football Pre-Season Rankings.docx'),
    '2025': Path(r'C:\Users\noahk\Downloads\2025 Fantasy Football Pre-Season Rankings.docx'),
    '2026': Path(r'C:\Users\noahk\Downloads\2026 Fantasy Football Pre-Season Rankings.docx'),
}
failed = False
for year, source in sources.items():
    output = (ROOT / 'pages' / f'{year}.html').read_text(encoding='utf-8')
    source_text = [p.text.strip() for p in Document(source).paragraphs if p.text.strip()]
    missing = [text for text in source_text if escape(text) not in output]
    image_count = len(Document(source).inline_shapes)
    output_images = output.count('<figure class="source-image">')
    print(f'{year}: {len(source_text) - len(missing)}/{len(source_text)} paragraphs preserved; {output_images}/{image_count} images preserved')
    if missing or image_count != output_images:
        failed = True
        for text in missing[:3]: print(f'  MISSING: {text[:100]}')
raise SystemExit(1 if failed else 0)
