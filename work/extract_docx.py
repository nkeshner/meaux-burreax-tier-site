from pathlib import Path
import sys
from docx import Document

def walk(parent, depth=0):
    for child in parent.element.body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            for paragraph in parent.paragraphs:
                if paragraph._p == child:
                    text = paragraph.text.strip()
                    if text:
                        print(f"P\t{paragraph.style.name}\t{text}")
                    break
        elif tag == 'tbl':
            for table in parent.tables:
                if table._tbl == child:
                    print("TABLE")
                    for row in table.rows:
                        print("ROW\t" + " | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
                    print("END_TABLE")
                    break

for value in sys.argv[1:]:
    path = Path(value)
    print(f"\n===== {path.name} =====")
    doc = Document(path)
    print(f"IMAGES\t{len(doc.inline_shapes)}")
    print(f"TABLES\t{len(doc.tables)}")
    walk(doc)
